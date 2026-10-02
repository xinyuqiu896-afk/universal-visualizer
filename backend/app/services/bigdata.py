"""大数据文件引擎：xlsx/csv 全量流式扫描 -> 真实聚合统计 + 明细 NDJSON 落盘 + 分页查询
设计要点：LLM 只见到聚合统计（KB 级），明细 17 万条通过 /api/records/{id}/data 分页供给，
从架构上杜绝「LLM 手写数据导致的缺失/编造」。
"""
import json, re
from collections import Counter
from pathlib import Path
from .. import config

BIGDATA_DIR = config.DATA_DIR / "bigdata"
BIGDATA_DIR.mkdir(parents=True, exist_ok=True)
BIGDATA_THRESHOLD = 1000          # 行数超过此值启用大文件模式
PAGE_DEFAULT_SIZE = 100

# 列语义识别规则（按表头关键字）
COL_RULES = {
    "name": ["企业名称", "名称", "公司", "单位"],
    "phone": ["手机号", "电话", "联系方式", "手机"],
    "carrier": ["运营商", "运营商归属"],
    "segment": ["号段"],
    "source": ["数据来源", "来源", "行业"],
}

def _detect_cols(headers):
    cols = {}
    for key, kws in COL_RULES.items():
        for i, h in enumerate(headers):
            if h and any(kw in str(h) for kw in kws):
                cols[key] = i
                break
    return cols

def scan_and_cache(file_id: str, path: str) -> dict:
    """全量扫描（read_only 流式，不占内存），统计聚合 + 明细写 NDJSON"""
    from openpyxl import load_workbook
    p = Path(path)
    ext = p.suffix.lower()

    ndjson_path = BIGDATA_DIR / f"{file_id}.ndjson"
    total = 0
    carriers, sources, segments = Counter(), Counter(), Counter()
    headers = []

    if ext in (".xlsx", ".xls"):
        wb = load_workbook(str(p), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        with open(ndjson_path, "w", encoding="utf-8") as f:
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                vals = ["" if c is None else str(c).strip() for c in row]
                if i == 0:
                    headers = vals
                    continue
                if not any(vals):
                    continue
                f.write(json.dumps(vals, ensure_ascii=False) + "\n")
                total += 1
        wb.close()
    elif ext == ".csv":
        import csv as csvmod
        with open(p, "r", encoding="utf-8-sig", errors="ignore", newline="") as fin, \
             open(ndjson_path, "w", encoding="utf-8") as f:
            reader = csvmod.reader(fin)
            for i, row in enumerate(reader):
                if i == 0:
                    headers = row
                    continue
                if not any(row):
                    continue
                f.write(json.dumps(row, ensure_ascii=False) + "\n")
                total += 1
    else:
        raise ValueError(f"大文件模式仅支持 xlsx/csv，当前 {ext}")

    # 表头已就绪后识别列语义，再对 NDJSON 二遍聚合（17 万行约数秒）
    cols = _detect_cols(headers)
    ci_carrier, ci_source, ci_seg = cols.get("carrier"), cols.get("source"), cols.get("segment")
    ci_phone = cols.get("phone")
    with open(ndjson_path, "r", encoding="utf-8") as f:
        for line in f:
            vals = json.loads(line)
            if ci_carrier is not None and ci_carrier < len(vals) and vals[ci_carrier]:
                carriers[vals[ci_carrier]] += 1
            if ci_source is not None and ci_source < len(vals) and vals[ci_source]:
                sources[vals[ci_source]] += 1
            if ci_seg is not None and ci_seg < len(vals) and vals[ci_seg]:
                segments[vals[ci_seg]] += 1
            elif ci_seg is None and ci_phone is not None and ci_phone < len(vals):
                phone = vals[ci_phone]
                if re.match(r"^1\d{10}$", phone):
                    segments[phone[:3]] += 1

    stats = {
        "total": total,
        "headers": headers,
        "col_map": cols,
        "carriers": dict(carriers.most_common(20)),
        "sources": dict(sources.most_common(30)),
        "segments": dict(segments.most_common(15)),
        "columns": [h for h in headers if h],
    }
    (BIGDATA_DIR / f"{file_id}.stats.json").write_text(
        json.dumps(stats, ensure_ascii=False), encoding="utf-8")
    stats["file_id"] = file_id
    return stats

def get_stats(file_id: str) -> dict | None:
    f = BIGDATA_DIR / f"{file_id}.stats.json"
    if not f.exists():
        return None
    return json.loads(f.read_text(encoding="utf-8"))

def load_page(file_id: str, page: int = 1, size: int = PAGE_DEFAULT_SIZE,
              carrier: str = "", source: str = "", keyword: str = "") -> dict:
    """明细分页查询（无过滤走行偏移快路径；有过滤全量流扫）"""
    ndjson_path = BIGDATA_DIR / f"{file_id}.ndjson"
    if not ndjson_path.exists():
        raise FileNotFoundError("该记录无大文件明细数据")
    stats = get_stats(file_id) or {}
    cols = stats.get("col_map", {})
    ci_carrier, ci_source, ci_name = cols.get("carrier"), cols.get("source"), cols.get("name")
    size = max(1, min(500, size))
    page = max(1, page)

    filtered = bool((carrier and carrier != "全部") or (source and source != "全部") or keyword.strip())
    rows, total = [], 0

    if not filtered:
        total = stats.get("total", 0)
        start = (page - 1) * size
        with open(ndjson_path, "r", encoding="utf-8") as f:
            for idx, line in enumerate(f):
                if idx < start:
                    continue
                rows.append(json.loads(line))
                if len(rows) >= size:
                    break
    else:
        kw = keyword.strip().lower()
        skip = (page - 1) * size
        with open(ndjson_path, "r", encoding="utf-8") as f:
            for line in f:
                vals = json.loads(line)
                if carrier and carrier != "全部" and (ci_carrier is None or vals[ci_carrier] != carrier):
                    continue
                if source and source != "全部" and (ci_source is None or vals[ci_source] != source):
                    continue
                if kw and not any(kw in str(v).lower() for v in vals):
                    continue
                total += 1
                if skip > 0:
                    skip -= 1
                    continue
                if len(rows) < size:
                    rows.append(vals)
    return {"total": total, "page": page, "size": size, "rows": rows,
            "headers": stats.get("headers", [])}
