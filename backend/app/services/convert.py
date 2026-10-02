"""格式转换引擎：真实转换，输出文件供下载"""
import csv
from pathlib import Path
from .. import config

async def convert(src: str, target_ext: str) -> dict:
    p = Path(src)
    ext = p.suffix.lower()
    out_name = p.stem + "_converted" + target_ext
    out = config.EXPORT_DIR / out_name

    if target_ext == ".txt":
        from .parser import parse_file
        text, _ = parse_file(src)
        out.write_text(text, encoding="utf-8")

    elif target_ext == ".docx":
        from docx import Document
        doc = Document()
        if ext in (".txt", ".md", ".json"):
            text = p.read_text(encoding="utf-8", errors="ignore")
            for line in text.splitlines():
                if line.strip():
                    (doc.add_heading(line.lstrip("#").strip(), level=1 if line.startswith("# ") else 2)
                     if line.startswith("#") else doc.add_paragraph(line))
        elif ext in (".csv", ".xlsx", ".xls"):
            rows = _read_table(src)
            if rows:
                t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                t.style = "Light Grid Accent 1"
                for i, r in enumerate(rows):
                    for j, c in enumerate(r):
                        t.cell(i, j).text = str(c)
        else:
            from .parser import parse_file
            for line in parse_file(src)["text"].splitlines():
                if line.strip():
                    doc.add_paragraph(line)
        doc.save(str(out))

    elif target_ext == ".pdf":
        # 纯文本/表格 -> PDF（用 reportlab 的替代：先转 docx 思路太重，这里用 fpdf2 风格的轻量实现）
        from fpdf import FPDF
        from .parser import parse_file
        text = parse_file(src)["text"]
        pdf = FPDF()
        pdf.add_page()
        # 使用 Windows 自带字体支持中文
        font_path = r"C:\Windows\Fonts\msyh.ttc"
        try:
            pdf.add_font("msyh", "", font_path)
            pdf.set_font("msyh", size=11)
        except Exception:
            pdf.set_font("helvetica", size=11)
        for line in text.splitlines():
            pdf.multi_cell(0, 6, line)
        pdf.output(str(out))

    elif target_ext in (".xlsx", ".csv"):
        rows = _read_table(src)
        if target_ext == ".csv":
            with open(out, "w", encoding="utf-8-sig", newline="") as f:
                csv.writer(f).writerows(rows)
        else:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "converted"
            for r in rows:
                ws.append(r)
            wb.save(str(out))
    else:
        raise ValueError(f"不支持转换到 {target_ext}")

    return {"file_name": out.name, "path": str(out), "size_kb": round(out.stat().st_size / 1024, 1)}

def _read_table(src) -> list:
    p = Path(src)
    ext = p.suffix.lower()
    if ext == ".csv":
        with open(p, "r", encoding="utf-8-sig", errors="ignore", newline="") as f:
            return [row for i, row in enumerate(csv.reader(f)) if i < 2000]
    elif ext in (".xlsx", ".xls"):
        from openpyxl import load_workbook
        wb = load_workbook(str(p), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = []
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            rows.append(["" if c is None else str(c) for c in r])
            if i >= 2000:
                break
        wb.close()
        return rows
    elif ext == ".docx":
        from docx import Document
        doc = Document(str(p))
        return [[c.text for c in r.cells] for t in doc.tables for r in t.rows]
    raise ValueError(f"格式 {ext} 无法转为表格")
